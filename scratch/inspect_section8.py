import docx
import sys
sys.stdout.reconfigure(encoding='utf-8')

doc = docx.Document(r'C:\Users\PRANAY KUMAR\Downloads\Review3_ACAFDM_Vehicular_Flooding (2).docx')

print("=== PARAGRAPHS 170 to 188 ===")
for i in range(170, min(188, len(doc.paragraphs))):
    p = doc.paragraphs[i]
    r_fonts = [r.font.name for r in p.runs if r.font and r.font.name]
    r_sizes = [r.font.size.pt for r in p.runs if r.font and r.font.size]
    style_name = p.style.name if p.style else 'None'
    print(f"P{i} | Style: {style_name} | Fonts: {r_fonts} | Sizes: {r_sizes}")
    print(f"    Text: {repr(p.text)}")

print("\n=== TABLES 14 to 17 ===")
for t_idx in [14, 15, 16, 17]:
    t = doc.tables[t_idx]
    print(f"\n--- Table {t_idx} (Style: {t.style.name if t.style else 'None'}) ---")
    for r_idx, row in enumerate(t.rows):
        cells = [c.text.strip().replace('\n', ' ') for c in row.cells]
        print(f"  Row {r_idx}: {cells}")

print("\n=== SECTION 8 SURROUNDING CONTEXT ===")
print("Paragraph right before Section 8:")
print(f"P170: {repr(doc.paragraphs[170].text)}")
print("Paragraph right after Section 8:")
print(f"P186: {repr(doc.paragraphs[186].text)}")
